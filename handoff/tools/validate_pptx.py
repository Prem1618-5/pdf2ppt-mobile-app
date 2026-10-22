"""Validate a generated deck against the two supported layouts.

    python validate_pptx.py deck.pptx --layout top|top-right [--no-blank] [--expect-questions N]

Exit code 0 = valid. Checks (these are the acceptance rules for the Android PPTX writer):
  * opens with python-pptx; slide size is 16:9 (13.333 x 7.5 in)
  * every picture lies inside the layout box (tolerance 0.02 in)
  * TOP        : box = x 0.15..13.183, y 0.15..3.60  (top half, full width)
  * TOP-RIGHT  : box = x 6.8165..13.183, y 0.15..3.60 (quadrant I: right of the vertical and
                 above the horizontal centre line, never crossing either)
  * with blank slides on, slides alternate question, blank, question, blank ...
Requires: pip install python-pptx
"""
import argparse
import sys
import zipfile

from pptx import Presentation
from pptx.util import Emu

SLIDE_W, SLIDE_H, INSET = 13.333, 7.5, 0.15
BOXES = {  # left, top, width, height (inches)
    "top": (INSET, INSET, SLIDE_W - 2 * INSET, SLIDE_H / 2 - 2 * INSET),
    "top-right": (SLIDE_W / 2 + INSET, INSET, SLIDE_W / 2 - 2 * INSET, SLIDE_H / 2 - 2 * INSET),
}
TOL = 0.02


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--layout", choices=BOXES, required=True)
    ap.add_argument("--no-blank", action="store_true", help="deck has no blank slide after each question")
    ap.add_argument("--expect-questions", type=int)
    a = ap.parse_args()

    errors = []
    with zipfile.ZipFile(a.deck) as z:
        names = set(z.namelist())
        for need in ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"):
            if need not in names:
                errors.append(f"missing part {need}")
        if z.testzip():
            errors.append("zip is corrupt")
    prs = Presentation(a.deck)
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    if abs(w - SLIDE_W) > 0.01 or abs(h - SLIDE_H) > 0.01:
        errors.append(f"slide size {w:.3f} x {h:.3f} in, expected {SLIDE_W} x {SLIDE_H}")
    L, T, BW, BH = BOXES[a.layout]
    questions = 0
    for i, slide in enumerate(prs.slides, 1):
        pics = [s for s in slide.shapes if s.shape_type == 13]
        texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        kind = "question" if (pics or texts) else "blank"
        if kind == "question":
            questions += 1
        if not a.no_blank:
            want = "question" if i % 2 == 1 else "blank"
            if kind != want:
                errors.append(f"slide {i}: expected {want}, found {kind}")
        for s in pics + texts:
            l, t = Emu(s.left).inches, Emu(s.top).inches
            r, b = l + Emu(s.width).inches, t + Emu(s.height).inches
            if l < L - TOL or t < T - TOL or r > L + BW + TOL or b > T + BH + TOL:
                errors.append(f"slide {i}: shape ({l:.2f},{t:.2f})-({r:.2f},{b:.2f}) outside {a.layout} box "
                              f"({L:.2f},{T:.2f})-({L + BW:.2f},{T + BH:.2f})")
            if b > SLIDE_H / 2 + TOL:
                errors.append(f"slide {i}: crosses the horizontal centre line")
            if a.layout == "top-right" and l < SLIDE_W / 2 - TOL:
                errors.append(f"slide {i}: crosses the vertical centre line")
    if a.expect_questions is not None and questions != a.expect_questions:
        errors.append(f"{questions} question slides, expected {a.expect_questions}")
    print(f"{a.deck}: {len(prs.slides)} slides, {questions} question slides, layout={a.layout}")
    for e in errors:
        print("  FAIL", e)
    print("OK" if not errors else f"{len(errors)} problem(s)")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
